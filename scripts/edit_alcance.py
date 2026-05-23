from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsmap
from copy import deepcopy
from lxml import etree

# Carregar documento original
doc = Document(
    "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/Artigo Revista Alcance.docx"
)

# namespaces
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Estratégia: modificar o documento original in-place
# 1. Adicionar autor após título ES (parágrafo 4)
# 2. Adicionar Resumen e Palabras clave após Keywords EN (parágrafo 19)

# Vamos usar uma abordagem direta de XML

# Obter o body do documento
body = doc._element.body

# Os parágrafos estão em body
# Parágrafo 4 = título ES
# Parágrafo 19 = Keywords EN

# Encontrar posições dos parágrafos no XML
paragraphs = body.findall(".//w:p", namespaces={"w": W})

print(f"Total de parágrafos XML: {len(paragraphs)}")

# Vamos examinar a estrutura
# Adicionar Resumen após Keywords EN
# Adicionar Palabras clave após Resumen
# Adicionar autor após título ES

# Criar elementos XML para Resumen e Palabras clave


def create_paragraph_element(text, bold=False, center=False, size=11, space_after=6):
    """Cria um elemento de parágrafo XML"""
    p = etree.Element(f"{{{W}}}p")

    pPr = etree.SubElement(p, f"{{{W}}}pPr")

    # spacing
    spacing = etree.SubElement(pPr, f"{{{W}}}spacing")
    spacing.set(f"{{{W}}}after", str(space_after * 20))  # 6pt = 120 twips

    if center:
        jc = etree.SubElement(pPr, f"{{{W}}}jc")
        jc.set(f"{{{W}}}val", "center")

    # run
    r = etree.SubElement(p, f"{{{W}}}r")
    rPr = etree.SubElement(r, f"{{{W}}}rPr")

    if bold:
        b = etree.SubElement(rPr, f"{{{W}}}b")

    sz = etree.SubElement(rPr, f"{{{W}}}sz")
    sz.set(f"{{{W}}}val", str(size * 2))  # half-points

    szCs = etree.SubElement(rPr, f"{{{W}}}szCs")
    szCs.set(f"{{{W}}}val", str(size * 2))

    rFonts = etree.SubElement(rPr, f"{{{W}}}rFonts")
    rFonts.set(f"{{{W}}}ascii", "Times New Roman")
    rFonts.set(f"{{{W}}}hAnsi", "Times New Roman")

    t = etree.SubElement(r, f"{{{W}}}t")
    t.text = text

    return p


def create_resumen_item(label, content):
    """Cria item do resumo estruturado"""
    p = etree.Element(f"{{{W}}}p")

    pPr = etree.SubElement(p, f"{{{W}}}pPr")
    spacing = etree.SubElement(pPr, f"{{{W}}}spacing")
    spacing.set(f"{{{W}}}after", str(6 * 20))

    # Label em negrito
    r1 = etree.SubElement(p, f"{{{W}}}r")
    rPr1 = etree.SubElement(r1, f"{{{W}}}rPr")
    b = etree.SubElement(rPr1, f"{{{W}}}b")
    sz = etree.SubElement(rPr1, f"{{{W}}}sz")
    sz.set(f"{{{W}}}val", "22")  # 11pt
    szCs = etree.SubElement(rPr1, f"{{{W}}}szCs")
    szCs.set(f"{{{W}}}val", "22")
    rFonts = etree.SubElement(rPr1, f"{{{W}}}rFonts")
    rFonts.set(f"{{{W}}}ascii", "Times New Roman")
    rFonts.set(f"{{{W}}}hAnsi", "Times New Roman")
    t1 = etree.SubElement(r1, f"{{{W}}}t")
    t1.text = label + " "

    # Conteúdo normal
    r2 = etree.SubElement(p, f"{{{W}}}r")
    rPr2 = etree.SubElement(r2, f"{{{W}}}rPr")
    sz2 = etree.SubElement(rPr2, f"{{{W}}}sz")
    sz2.set(f"{{{W}}}val", "22")
    szCs2 = etree.SubElement(rPr2, f"{{{W}}}szCs")
    szCs2.set(f"{{{W}}}val", "22")
    rFonts2 = etree.SubElement(rPr2, f"{{{W}}}rFonts")
    rFonts2.set(f"{{{W}}}ascii", "Times New Roman")
    rFonts2.set(f"{{{W}}}hAnsi", "Times New Roman")
    t2 = etree.SubElement(r2, f"{{{W}}}t")
    t2.text = content

    return p


# Encontrar parágrafo 4 (título ES) e 19 (Keywords EN)
# No XML, cada parágrafo é um elemento w:p

# Criar elemento de AUTOR
autor_para = create_paragraph_element(
    "Renato de Oliveira Rosa", center=True, space_after=12
)

# Criar elementos de RESUMEN
resumen_header = create_paragraph_element(
    "Resumen", bold=True, center=True, space_after=6
)

resumen_items_data = [
    (
        "Objetivo:",
        "Analizar críticamente si las declaraciones de ejecutivos que atribuyen despidos en masa a la implementación de inteligencia artificial encuentran respaldo en la realidad operacional pública o constituyen narrativa estratégica para ocultar motivaciones gerenciales subyacentes.",
    ),
    (
        "Diseño / metodología / enfoque:",
        "La investigación se caracteriza como cualitativa y se basa en análisis de contenido. Para ello, se utilizó como corpus documental las declaraciones y los estados financieros de seis empresas de tecnología (Salesforce, Intuit, Dropbox, Block, Cisco e IBM) que anuncieron reducciones de personal entre 2023 y 2026.",
    ),
    (
        "Resultados:",
        "Ninguna de las organizaciones disponibilizó públicamente métricas verificables de productividad que sustentaran la redundancia de trabajadores por la tecnología. En cambio, se observó rentabilidad, historial de sobrecontratación pandémica, lenguaje de determinismo tecnológico e inversiones simultáneas en la propia herramienta. De esta forma, se caracterizó el argumento tecnológico como pretexto discursivo.",
    ),
    (
        "Limitaciones / implicaciones de la investigación:",
        "La muestra se limitó al sector de tecnología estadounidense y a fuentes públicas. El fenómeno investigado es reciente y aún en evolución, lo que limita la generalización de los hallazgos.",
    ),
    (
        "Implicaciones prácticas:",
        "El trabajo proporciona un modelo analítico para que inversionistas, reguladores y trabajadores evalúen alegaciones corporativas. Con ese propósito, se possibilita la identificación de inconsistencias entre discursos de automatización y la capacidad financiera real de las organizaciones.",
    ),
    (
        "Implicaciones sociales:",
        "El estudio ofrece insumos para que la sociedad evalúe críticamente las alegaciones de empresas sobre la relación entre implementación de IA y despidos, contribuyendo a la transparencia en el mercado de trabajo.",
    ),
    (
        "Implicaciones teóricas:",
        "El estudio demuestra la persistencia de la Paradoja de Solow a nivel de la firma. De forma complementaria, se aplica la Teoría de la Agencia para explicar la externalización de responsabilidades ejecutivas por medio de narrativas de innovación.",
    ),
    (
        "Originalidad / valor:",
        "El estudio llena vacíos en la literatura al cruzar marcos de agencia, productividad y recursos para demostrar que la narrativa de IA funciona como cortina de humo para despidos motivados por otros factores.",
    ),
]

resumen_items = [
    create_resumen_item(label, content) for label, content in resumen_items_data
]

# Criar Palabras clave
palabras_para = create_paragraph_element(
    "Palabras clave: Paradoja de la productividad; Teoría de la agencia; Despidos corporativos; Inteligencia artificial; AI scapegoating; Discurso organizacional",
    bold=True,
    space_after=18,
)

# Agora inserir no XML
# Encontrar parágrafo 4 (título ES) - índice 4 no XML
# e parágrafo 19 (Keywords EN) - índice 19 no XML

# Iterar sobre os elementos do body para encontrar as posições
# No body, elementos podem ser: w:p (parágrafos), w:tbl (tabelas), etc.

elements = list(body)
print(f"Total de elementos no body: {len(elements)}")

# Encontrar índice do parágrafo 4 (título ES) e 19 (Keywords EN)
para_count = 0
idx_titulo_es = None
idx_keywords_en = None

for i, elem in enumerate(elements):
    if elem.tag == f"{{{W}}}p":
        if para_count == 4:
            idx_titulo_es = i
        if para_count == 19:
            idx_keywords_en = i
        para_count += 1

print(f"Índice do título ES no XML: {idx_titulo_es}")
print(f"Índice do Keywords EN no XML: {idx_keywords_en}")

# Inserir AUTOR após título ES (índice 4)
if idx_titulo_es is not None:
    autor_para_copy = deepcopy(autor_para)
    elements.insert(idx_titulo_es + 1, autor_para_copy)
    print("Autor inserido")

# Recalcular índices após inserção
idx_keywords_en += 1  # Shifted by 1

# Inserir Resumen, Items e Palabras clave após Keywords EN
if idx_keywords_en is not None:
    insert_pos = idx_keywords_en + 1

    # Inserir Resumen header
    resumen_header_copy = deepcopy(resumen_header)
    elements.insert(insert_pos, resumen_header_copy)
    insert_pos += 1

    # Inserir items
    for item in resumen_items:
        item_copy = deepcopy(item)
        elements.insert(insert_pos, item_copy)
        insert_pos += 1

    # Inserir Palabras clave
    palabras_copy = deepcopy(palabras_para)
    elements.insert(insert_pos, palabras_copy)
    print("Resumen e Palabras clave inseridos")

# Atualizar o body com os elementos modificados
# Primeiro, limpar o body
for elem in body:
    body.remove(elem)

# Re-adicionar todos os elementos
for elem in elements:
    body.append(elem)

# Salvar
output_path = "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/Artigo Revista Alcance_CORRIGIDO.docx"
doc.save(output_path)

print(f"\nDocumento salvo em: {output_path}")

# Verificação
verify_doc = Document(output_path)
full_text = " ".join([p.text for p in verify_doc.paragraphs])
word_count = len(full_text.split())

print(f"Palavras: {word_count}")
print(f"Parágrafos: {len(verify_doc.paragraphs)}")
print(f"Tabelas: {len(verify_doc.tables)}")

print("\n=== PRIMEIROS 45 PARÁGRAFOS ===")
for i, para in enumerate(verify_doc.paragraphs[:45]):
    text = para.text[:65] if para.text else "[vazio]"
    print(f"{i}: {text}")
