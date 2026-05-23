from docx import Document
from docx.shared import Inches, Pt
import os

BASE = "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA"
FILES = [
    os.path.join(BASE, "Artigo_Alcance_FINAL.docx"),
    os.path.join(BASE, "Manuscrito_Alcance_Anônimo.docx"),
    os.path.join(BASE, "Manuscrito_Alcance_ComAutores.docx"),
]

ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

for path in FILES:
    if not os.path.exists(path):
        print(f"Arquivo não encontrado, pulando: {path}")
        continue
    print(f"Processando: {path}")
    doc = Document(path)

    # helper to check if paragraph in table
    def para_in_table(par):
        try:
            ancestors = par._p.xpath('ancestor::w:tc', namespaces=ns)
            return len(ancestors) > 0
        except Exception:
            return False

    # remove first-line indent for all paragraphs
    for p in doc.paragraphs:
        try:
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.left_indent = Inches(0)
        except Exception:
            pass

    # also ensure paragraphs inside table cells have no indent
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    try:
                        p.paragraph_format.first_line_indent = Inches(0)
                        p.paragraph_format.left_indent = Inches(0)
                    except Exception:
                        pass

    doc.save(path)
    print(f"Salvo: {path}\n")

print('Concluído: recuos removidos nos arquivos disponíveis.')
