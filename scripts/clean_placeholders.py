from docx import Document
import re
from pathlib import Path

infile = Path(r"C:/Users/Renato/Documents/Artigos 2026/Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/Artigo_Alcance_FINAL_SUBMIT.docx")
outfile = infile.with_name(infile.stem + "_CLEAN" + infile.suffix)
report = Path(r"C:/Users/Renato/Documents/Artigos 2026/Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/clean_placeholders_report.txt")

if not infile.exists():
    print(f"Input file not found: {infile}")
    raise SystemExit(1)

doc = Document(infile)

# Detect start of references section by common headings or first reference-like line
ref_start_idx = None
ref_heading_pattern = re.compile(r'^(references|referências|referencias)$', re.I)
ref_like_pattern = re.compile(r'^[A-Z][^\n]+,\s+[A-Z]?[\.\w\- ]*\(\d{4}\)')

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if ref_start_idx is None and (ref_heading_pattern.match(t) or ref_like_pattern.match(t)):
        ref_start_idx = i
        break

if ref_start_idx is None:
    # Fallback: assume last third of doc is references
    ref_start_idx = int(len(doc.paragraphs) * 2 / 3)

placeholder_pattern = re.compile(r'\[(?:VERIFICAR|AUTOMATIC|AUTOMATIC PLACEHOLDER|AUTOMATIC_PLACEHOLDER|VERIFICAR:).*?\]', re.I)
removed = []

# Helper to remove paragraph element
def remove_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    try:
        paragraph._p = paragraph._element = None
    except Exception:
        pass

# Iterate over a snapshot to avoid mutation issues
for idx, p in enumerate(list(doc.paragraphs)):
    t = p.text.strip()
    if not t:
        continue
    if idx < ref_start_idx and placeholder_pattern.search(t):
        removed.append((idx, t))
        remove_paragraph(p)

# Save cleaned file and report
doc.save(outfile)

with report.open('w', encoding='utf-8') as f:
    f.write(f"Input: {infile}\nOutput: {outfile}\nRemoved paragraphs: {len(removed)}\n\n")
    for i, txt in removed:
        f.write(f"[{i}] {txt}\n")

print(f"Done. Removed {len(removed)} paragraphs. Clean file: {outfile}")
print(f"Report: {report}")
