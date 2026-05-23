from docx import Document
from docx.shared import Pt
import os

BASE = "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA"
COMAUTH = os.path.join(BASE, "Manuscrito_Alcance_ComAutores.docx")
LETTER = os.path.join(BASE, "Carta_Editor.docx")

author_name = 'Renato de Oliveira Rosa'
affiliation = 'Fucape Business School'
contact_email = ''

# Update manuscript with authors
if os.path.exists(COMAUTH):
    doc = Document(COMAUTH)
    # Insert after the three title paragraphs (they were added by generator)
    insert_idx = 3
    p = doc.paragraphs[insert_idx].insert_paragraph_before()
    p.alignment = 1
    run = p.add_run(author_name + '\n')
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'
    run2 = p.add_run(affiliation)
    run2.font.size = Pt(11)
    run2.font.name = 'Times New Roman'
    doc.save(COMAUTH)
    print('Autores inseridos em:', COMAUTH)
else:
    print('Arquivo com autores não encontrado:', COMAUTH)

# Update carta ao editor
if os.path.exists(LETTER):
    doc = Document(LETTER)
    # Insert at top
    p = doc.paragraphs[0]
    p.insert_paragraph_before(author_name)
    p.insert_paragraph_before(affiliation)
    doc.save(LETTER)
    print('Carta atualizada em:', LETTER)
else:
    print('Carta ao editor não encontrada:', LETTER)
