from docx import Document
import os

path = r'C:\Users\Renato\Documents\Artigos 2026\Prontos-Submissao\Artigo - Paradoxo da Produtividade IA\Artigo_Alcance_FINAL.docx'
doc = Document(path)
print('paragraphs', len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs[:80]):
    sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None
    sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None
    fi = p.paragraph_format.first_line_indent.pt if p.paragraph_format.first_line_indent else None
    al = p.alignment
    print(i, repr(p.text), 'align', al, 'sb', sb, 'sa', sa, 'fi', fi)
print('tables', len(doc.tables))
for ti, t in enumerate(doc.tables):
    print('table', ti, 'rows', len(t.rows), 'cols', len(t.columns))
print('inline drawings', sum(len(p._p.xpath('.//w:drawing')) for p in doc.paragraphs))
