from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Carregar documento original
original_doc = Document(
    "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/Artigo Revista Alcance -.docx"
)

# Criar novo documento
new_doc = Document()

# Configurar margens
sections = new_doc.sections
for section in sections:
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)


def add_formatted_paragraph(
    doc,
    text,
    bold=False,
    size=11,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line=True,
    space_after=6,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return p


# 1. TÍTULO EM PORTUGUÊS
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run(
    "O PARADOXO DA PRODUTIVIDADE DA IA: ANÁLISE CRÍTICA DO DISCURSO CORPORATIVO SOBRE DEMISSÕES TECNOLÓGICAS"
)
run.bold = True
run.font.size = Pt(14)
run.font.name = "Times New Roman"

# 2. TÍTULO EM INGLÊS
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run(
    "THE AI PRODUCTIVITY PARADOX: CRITICAL ANALYSIS OF CORPORATE DISCOURSE ON TECHNOLOGICAL LAYOFFS"
)
run.bold = True
run.font.size = Pt(14)
run.font.name = "Times New Roman"

# 3. TÍTULO EM ESPANHOL
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
run = p.add_run(
    "EL PARADOJA DE LA PRODUCTIVIDAD DE LA IA: ANÁLISIS CRÍTICO DEL DISCURSO CORPORATIVO SOBRE DESPIDOS TECNOLÓGICOS"
)
run.bold = True
run.font.size = Pt(14)
run.font.name = "Times New Roman"

# 4. AUTORES
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
run = p.add_run("Renato de Oliveira Rosa¹; Coautor(es)²")
run.font.size = Pt(11)
run.font.name = "Times New Roman"

# 5. RESUMO - PORTUGUÊS
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run("Resumo")
run.bold = True
run.font.size = Pt(12)
run.font.name = "Times New Roman"

resumo_items = [
    (
        "Objetivo:",
        "Analisar criticamente se as declarações de executivos que atribuem demissões em massa à implementação de inteligência artificial encontram respaldo na realidade operacional pública ou se constituem narrativa estratégica para ocultar motivações gerenciais subjacentes.",
    ),
    (
        "Design / metodologia / abordagem:",
        "A pesquisa caracteriza-se como qualitativa e baseia-se em análise de conteúdo. Para tanto, utilizou-se como corpus documental as declarações e os demonstrativos financeiros de seis empresas de tecnologia (Salesforce, Intuit, Dropbox, Block, Cisco e IBM) que anunciaram reduções de pessoal entre 2023 e 2026.",
    ),
    (
        "Resultados:",
        "Nenhuma das organizações disponibilizou publicamente métricas verificáveis de produtividade que sustentassem a redundância de trabalhadores pela tecnologia. Em contrapartida, observou-se lucratividade, histórico de sobrecontratação pandêmica, linguagem de determinismo tecnológico e investimentos simultâneos na própria ferramenta. Dessa forma, caracterizou-se o argumento tecnológico como pretexto discursivo.",
    ),
    (
        "Limitações / implicações da pesquisa:",
        "A amostra foi limitada ao setor de tecnologia norte-americano e a fontes públicas. O fenômeno investigado é recente e ainda em evolução, o que limita a generalização dos achados.",
    ),
    (
        "Implicações práticas:",
        "O trabalho fornece um modelo analítico para que investidores, reguladores e trabalhadores avaliem alegações corporativas. Com esse propósito, possibilita-se a identificação de inconsistências entre discursos de automação e a capacidade financeira real das organizações.",
    ),
    (
        "Implicações sociais:",
        "O estudo oferece subsídios para que a sociedade avalie criticamente alegações de empresas sobre a relação entre implementação de IA e demissões, contribuindo para transparência no mercado de trabalho.",
    ),
    (
        "Implicações teóricas:",
        "O estudo demonstra a persistência do Paradoxo de Solow no nível da firma. De forma complementar, aplica-se a Teoria da Agência para explicar a externalização de responsabilidades executivas por meio de narrativas de inovação.",
    ),
    (
        "Originalidade / valor:",
        "O estudo preenche lacunas na literatura ao cruzar referenciais de agência, produtividade e recursos para demonstrar que a narrativa de IA funciona como cortina de fumaça para demissões motivadas por outros fatores.",
    ),
]

for label, content in resumo_items:
    p = new_doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run_label = p.add_run(label + " ")
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_label.font.name = "Times New Roman"
    run_content = p.add_run(content)
    run_content.font.size = Pt(11)
    run_content.font.name = "Times New Roman"

# Palavras-chave em português
p = new_doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
run = p.add_run("Palavras-chave: ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Times New Roman"
run = p.add_run(
    "Paradoxo da produtividade; Teoria da agência; Demissões corporativas; Inteligência artificial; AI washing"
)
run.font.size = Pt(11)
run.font.name = "Times New Roman"

# ABSTRACT - INGLÊS
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run("Abstract")
run.bold = True
run.font.size = Pt(12)
run.font.name = "Times New Roman"

abstract_items_en = [
    (
        "Purpose:",
        "To critically analyze whether executive statements attributing mass layoffs to artificial intelligence implementation find support in public operational reality or constitute strategic narratives to conceal underlying managerial motivations.",
    ),
    (
        "Design/methodology/approach:",
        "This qualitative research was based on documentary content analysis. The corpus comprised statements and financial reports from six technology companies (Salesforce, Intuit, Dropbox, Block, Cisco, and IBM) that announced layoffs between 2023 and 2026.",
    ),
    (
        "Findings:",
        "No organization publicly provided verifiable productivity metrics supporting human redundancy through technology. In contrast, profitability, pandemic-era over-hiring history, technological determinism language, and simultaneous investments in the same tool were observed. Thus, the technological argument was characterized as discursive pretext.",
    ),
    (
        "Research limitations/implications:",
        "The sample was limited to the US technology sector and public sources. The investigated phenomenon is recent and still evolving, limiting the generalization of findings.",
    ),
    (
        "Practical implications:",
        "The paper provides an analytical model for investors, regulators, and workers to evaluate corporate allegations. This enables the identification of inconsistencies between automation discourses and the organizations' real financial capacity.",
    ),
    (
        "Social implications:",
        "The study offers inputs for society to critically evaluate companies' claims about the relationship between AI implementation and layoffs, contributing to labor market transparency.",
    ),
    (
        "Theoretical implications:",
        "The study demonstrates the persistence of Solow's Paradox at the firm level. Complementarily, it applies Agency Theory to explain the externalization of executive responsibilities through innovation narratives.",
    ),
    (
        "Originality/value:",
        "This article fills literature gaps by crossing agency, productivity, and resource frameworks to demonstrate that the AI narrative functions as a smoke screen for layoffs motivated by other factors.",
    ),
]

for label, content in abstract_items_en:
    p = new_doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run_label = p.add_run(label + " ")
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_label.font.name = "Times New Roman"
    run_content = p.add_run(content)
    run_content.font.size = Pt(11)
    run_content.font.name = "Times New Roman"

# Keywords em inglês
p = new_doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
run = p.add_run("Keywords: ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Times New Roman"
run = p.add_run(
    "Productivity paradox; Agency theory; Corporate layoffs; Artificial intelligence; AI scapegoating; Organizational discourse"
)
run.font.size = Pt(11)
run.font.name = "Times New Roman"

# RESUMEN - ESPANHOL
p = new_doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run("Resumen")
run.bold = True
run.font.size = Pt(12)
run.font.name = "Times New Roman"

resumen_items = [
    (
        "Objetivo:",
        "Analizar críticamente si las declaraciones de ejecutivos que atribuyen despidos en masa a la implementación de inteligencia artificial encuentran respaldo en la realidad operacional pública o constituyen narrativa estratégica para ocultar motivaciones gerenciales subyacentes.",
    ),
    (
        "Diseño / metodología / enfoque:",
        "La investigación se caracteriza como cualitativa y se basa en análisis de contenido. Para ello, se utilizó como corpus documental las declaraciones y los estados financieros de seis empresas de tecnología (Salesforce, Intuit, Dropbox, Block, Cisco e IBM) que anuncieron reducciones de personal entre 2023 y 2026.",
    ),
    (
        "Resultados:",
        "Ninguna de las organizaciones disponibilizó públicamente métricas verificables de productividad que sustentaran la redundancia de trabajadores por la tecnología. En cambio, se observó rentabilidad, historial de sobrecontratación pandémica, lenguaje de determinismo tecnológico e inversiones simultáneas en la propia herramienta. De esta forma, se caracterizó el argumento tecnológico como pretexto discursivo.",
    ),
    (
        "Limitaciones / implicaciones de la investigación:",
        "La muestra se limitó al sector de tecnología estadounidense y a fuentes públicas. El fenómeno investigado es reciente y aún en evolución, lo que limita la generalización de los hallazgos.",
    ),
    (
        "Implicaciones prácticas:",
        "El trabajo proporciona un modelo analítico para que inversionistas, reguladores y trabajadores evalúen alegaciones corporativas. Con ese propósito, se posibilita la identificación de inconsistencias entre discursos de automatización y la capacidad financiera real de las organizaciones.",
    ),
    (
        "Implicaciones sociales:",
        "El estudio ofrece insumos para que la sociedad evalúe críticamente las alegaciones de empresas sobre la relación entre implementación de IA y despidos, contribuyendo a la transparencia en el mercado de trabajo.",
    ),
    (
        "Implicaciones teóricas:",
        "El estudio demuestra la persistencia de la Paradoja de Solow a nivel de la firma. De forma complementaria, se aplica la Teoría de la Agencia para explicar la externalización de responsabilidades ejecutivas por medio de narrativas de innovación.",
    ),
    (
        "Originalidad / valor:",
        "El estudio llena vacíos en la literatura al cruzar marcos de agencia, productividad y recursos para demostrar que la narrativa de IA funciona como cortina de humo para despidos motivados por otros factores.",
    ),
]

for label, content in resumen_items:
    p = new_doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run_label = p.add_run(label + " ")
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_label.font.name = "Times New Roman"
    run_content = p.add_run(content)
    run_content.font.size = Pt(11)
    run_content.font.name = "Times New Roman"

# Palabras clave em espanhol
p = new_doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
run = p.add_run("Palabras clave: ")
run.bold = True
run.font.size = Pt(11)
run.font.name = "Times New Roman"
run = p.add_run(
    "Paradoja de la productividad; Teoría de la agencia; Despidos corporativos; Inteligencia artificial; AI scapegoating; Discurso organizacional"
)
run.font.size = Pt(11)
run.font.name = "Times New Roman"

# Copiar conteúdo do artigo original a partir da Introdução
intro_idx = None
for i, para in enumerate(original_doc.paragraphs):
    if "1. Introdução" in para.text or "INTRODUÇÃO" in para.text.upper():
        intro_idx = i
        break

print(f"Introdução encontrada no índice: {intro_idx}")

if intro_idx:
    for i in range(intro_idx, len(original_doc.paragraphs)):
        para = original_doc.paragraphs[i]
        text = para.text.strip()
        if not text:
            continue

        new_para = new_doc.add_paragraph()
        new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        new_para.paragraph_format.space_after = Pt(6)
        new_para.paragraph_format.first_line_indent = Inches(0.5)

        # Preservar negrito se existir
        is_bold = any(run.bold for run in para.runs)
        for run in para.runs:
            new_run = new_para.add_run(run.text)
            new_run.font.name = "Times New Roman"
            new_run.font.size = Pt(11)
            new_run.bold = run.bold

# Copiar tabelas
for table in original_doc.tables:
    new_table = new_doc.add_table(rows=len(table.rows), cols=len(table.columns))
    new_table.style = "Table Grid"

    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            new_table.rows[i].cells[j].text = cell.text

# Salvar documento corrigido
output_path = "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA/Artigo_Alcance_Corrigido.docx"
new_doc.save(output_path)
print(f"Documento corrigido salvo em: {output_path}")

# Contar palavras
full_text = " ".join([p.text for p in new_doc.paragraphs])
word_count = len(full_text.split())
print(f"Total de palavras no documento corrigido: {word_count}")
