import os
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

BASE = Path('Prontos-Submissao/Artigo - Paradoxo da Produtividade IA')
FINAL_FILES = [
    BASE / 'Artigo_Alcance_FINAL.docx',
    BASE / 'Manuscrito_Alcance_Anônimo.docx',
]
SRC_DOCX = BASE / 'paradoxo_produtividade_ia_v1.docx'

MEDIA_FILES = {
    'figure_1': 'word/media/image3.png',
    'figure_2': 'word/media/image1.png',
}

OUTPUT_SUFFIX = '_FIXED.docx'


def extract_media_bytes(docx_path: Path, media_path: str) -> bytes:
    with ZipFile(docx_path, 'r') as z:
        return z.read(media_path)


def find_paragraph_index(doc: Document, needle: str):
    for idx, p in enumerate(doc.paragraphs):
        if needle in p.text:
            return idx
    return None


def insert_paragraph_after(doc: Document, target_idx: int, text: str = ''):
    target = doc.paragraphs[target_idx]
    new_element = OxmlElement('w:p')
    target._element.addnext(new_element)
    new_para = Paragraph(new_element, doc)
    if text:
        new_para.add_run(text)
    return new_para


def fix_references(doc: Document):
    ref_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip().lower().startswith('referências') or p.text.strip().lower().startswith('referencias'):
            ref_idx = idx
            break
    if ref_idx is None:
        print('  ⚠️ Não achei o cabeçalho de referências no documento.', doc)
        return

    for p in doc.paragraphs[ref_idx + 1:]:
        if not p.text.strip():
            continue
        if p.text.strip().lower().startswith('referências') or p.text.strip().lower().startswith('referencias'):
            continue
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def remove_drawings_from_paragraph(paragraph):
    removed = False
    for node in list(paragraph._p.iter()):
        if node.tag.endswith('drawing'):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
                removed = True
    return removed


def add_figure_after_caption(doc: Document, caption_text: str, image_bytes: bytes, width=Inches(5.8)):
    caption_idx = find_paragraph_index(doc, caption_text)
    if caption_idx is None:
        print(f'  ⚠️ Não achei o texto de legenda "{caption_text}" no documento.')
        return False

    # insert a blank paragraph after the caption paragraph, then add image to it
    image_para = insert_paragraph_after(doc, caption_idx)
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = image_para.add_run()
    run.add_picture(BytesIO(image_bytes), width=width)
    image_para.paragraph_format.space_after = Pt(6)
    image_para.paragraph_format.space_before = Pt(0)
    return True


def fix_document(path: Path):
    print('Processando:', path.name)
    doc = Document(path)

    # Fix and restore references formatting
    fix_references(doc)

    # Remove broken drawing from the current figure-1 caption area, then reinsert the correct image.
    image3_bytes = extract_media_bytes(SRC_DOCX, MEDIA_FILES['figure_1'])
    fig1_caption = 'Relação Percentual entre Redução Funcional e Faturamento Corporativo'
    fig1_idx = find_paragraph_index(doc, fig1_caption)
    if fig1_idx is not None and fig1_idx + 1 < len(doc.paragraphs):
        removed = remove_drawings_from_paragraph(doc.paragraphs[fig1_idx + 1])
        if removed:
            print('  ✅ Removido desenho quebrado da área da Figura 1.')
    if add_figure_after_caption(doc, fig1_caption, image3_bytes):
        print('  ✅ Figura 1 restaurada.')
    else:
        print('  ⚠️ Não consegui inserir a Figura 1 no documento.')

    # Insert the missing figure after the caption for Figura 2
    image1_bytes = extract_media_bytes(SRC_DOCX, MEDIA_FILES['figure_2'])
    if add_figure_after_caption(doc, 'Evolução do Quadro de Colaboradores: Pré-pandemia, Pico e Pós-demissão', image1_bytes):
        print('  ✅ Figura 2 restaurada.')
    else:
        print('  ⚠️ Não consegui inserir a Figura 2 no documento.')

    out_path = path.with_name(path.stem + OUTPUT_SUFFIX)
    doc.save(out_path)
    print('  Salvou:', out_path.name)
    return out_path


def count_words(doc: Document):
    import re
    total = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        total += len(re.findall(r'\w+', text, flags=re.UNICODE))
    return total


if __name__ == '__main__':
    for path in FINAL_FILES:
        if not path.exists():
            print('Arquivo não encontrado:', path)
            continue
        out_path = fix_document(path)
        doc = Document(out_path)
        print('  Palavras (após correção):', count_words(doc))
        print()
