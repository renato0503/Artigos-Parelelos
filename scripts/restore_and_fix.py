import shutil
import os
from docx import Document
from docx.shared import Pt, Inches

BASE = "Prontos-Submissao/Artigo - Paradoxo da Produtividade IA"
SRC = os.path.join(BASE, "Artigo Revista Alcance.docx")
FINAL = os.path.join(BASE, "Artigo_Alcance_FINAL.docx")
ANON = os.path.join(BASE, "Manuscrito_Alcance_Anônimo.docx")

if not os.path.exists(SRC):
    print('Arquivo fonte não encontrado:', SRC)
    raise SystemExit(1)

print('Sobrescrevendo FINAL com o arquivo fonte (preserva imagens)...')
shutil.copyfile(SRC, FINAL)

print('Carregando FINAL para ajustes...')
doc = Document(FINAL)

# helper to check if paragraph in table
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def para_in_table(par):
    try:
        ancestors = par._p.xpath('ancestor::w:tc', namespaces=ns)
        return len(ancestors) > 0
    except Exception:
        return False

print('Ajustando parágrafos: aplicar espaçamento apenas fora das tabelas...')
for p in doc.paragraphs:
    if para_in_table(p):
        try:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0)
        except Exception:
            pass
    else:
        try:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Inches(0.5)
        except Exception:
            pass

print('Ajustando parágrafos dentro das células das tabelas (garantir 0 spacing)...')
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                try:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = Inches(0)
                except Exception:
                    pass

print('Salvando FINAL corrigido (sobrescrevendo)...')
doc.save(FINAL)

print('Atualizando Manuscrito anônimo no lugar (removendo bloco de autores)...')
anon = Document(FINAL)
stop_idx = None
for i, p in enumerate(anon.paragraphs[:80]):
    t = p.text.strip().lower()
    if t.startswith('resumo') or t.startswith('abstract'):
        stop_idx = i
        break
paras_to_remove = []
if stop_idx and stop_idx > 3:
    for i in range(3, stop_idx):
        paras_to_remove.append(anon.paragraphs[i])
for p in paras_to_remove:
    try:
        p._element.getparent().remove(p._element)
    except Exception:
        pass
anon.core_properties.author = ''
anon.save(ANON)

print('Concluído. Arquivos atualizados:')
print(' -', FINAL)
print(' -', ANON)
